"""Build versioned CRM Word deliverables from the canonical Spanish topics.

python-docx is an authoring-only dependency: it provides native tables, numbering,
relationships and image preservation that the baseline PowerShell writer lacks.
Install it in a temporary virtual environment; the CRM runtime does not use it.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree


def sha256(data: bytes) -> str:
    """Return a stable digest for provenance and preservation checks."""
    return hashlib.sha256(data).hexdigest().upper()


def element(tag: str, **attributes):
    """Create one namespace-qualified Word element."""
    node = OxmlElement(tag)
    for key, value in attributes.items():
        node.set(qn('w:' + key), str(value))
    return node


class ManualWriter:
    """Render one publication with shared native Word styles and controls."""

    def __init__(self, config: dict, revision: str, title: str):
        self.config = config
        self.revision = revision
        self.publication_date = datetime.strptime(config['publicationDate'], '%Y-%m-%d').replace(tzinfo=timezone.utc)
        self.doc = Document()
        self.entries = []
        self.chapter = 0
        self.topic = 0
        self.figure = 0
        self.media = []
        self.sources = {}
        self.ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        self._configure(title)

    def _configure(self, title: str):
        """Set A4 layout, accessible heading hierarchy and corporate typography."""
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
        section.left_margin, section.right_margin = Cm(2), Cm(2)
        section.header_distance = section.footer_distance = Cm(0.8)
        normal = self.doc.styles['Normal']
        normal.font.name = 'Arial'
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGBColor.from_string('243247')
        normal.paragraph_format.space_after = Pt(5)
        normal.paragraph_format.line_spacing = 1.08
        normal.paragraph_format.widow_control = True
        for style_name, size, color in [('Title', 30, '00296B'), ('Subtitle', 13, '244C82'),
                                       ('Heading 1', 18, '00296B'), ('Heading 2', 13, '244C82'),
                                       ('Heading 3', 10.5, '00296B'), ('Heading 4', 10.5, '244C82')]:
            style = self.doc.styles[style_name]
            style.font.name = 'Arial'
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string(color)
            style.font.bold = style_name != 'Subtitle'
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(13 if style_name == 'Heading 1' else 8)
            style.paragraph_format.space_after = Pt(5)
        self.doc.styles['Caption'].font.name = 'Arial'
        self.doc.styles['Caption'].font.size = Pt(9)
        self.doc.styles['Caption'].font.color.rgb = RGBColor.from_string('526176')
        settings = self.doc.settings.element
        settings.append(element('w:updateFields', val='true'))
        for style in [normal, *(self.doc.styles[f'Heading {level}'] for level in range(1, 5))]:
            style.element.get_or_add_rPr().append(element('w:lang', val='es-ES'))
        self.heading_num = self._number_definition('heading')
        core = self.doc.core_properties
        core.title, core.author = title, 'INSERTEC'
        core.subject = 'Documentacion de uso de la aplicacion web CRM'
        core.language = 'es-ES'
        core.version = self.config['version'] + '-' + self.revision
        core.created = core.modified = self.publication_date
        header = section.header.paragraphs[0]
        header.text = 'INSERTEC  |  APLICACION WEB CRM'
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor.from_string('64748B')
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run(f"Edicion {self.config['version']} · {self.revision}  |  ")
        self._field(footer, 'PAGE', '1')
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string('64748B')

    def _number_definition(self, kind: str) -> int:
        """Create a native list instance; ordered procedures restart at one."""
        root = self.doc.part.numbering_part.element
        abstract_ids = [int(n.get(qn('w:abstractNumId'))) for n in root.findall(qn('w:abstractNum'))]
        num_ids = [int(n.get(qn('w:numId'))) for n in root.findall(qn('w:num'))]
        abstract_id, num_id = max(abstract_ids, default=0) + 1, max(num_ids, default=0) + 1
        abstract = element('w:abstractNum', abstractNumId=abstract_id)
        abstract.append(element('w:multiLevelType', val='multilevel' if kind == 'heading' else 'singleLevel'))
        for level in range(2 if kind == 'heading' else 1):
            lvl = element('w:lvl', ilvl=level)
            lvl.append(element('w:start', val=1))
            lvl.append(element('w:numFmt', val='bullet' if kind == 'bullet' else 'decimal'))
            if kind == 'heading':
                lvl.append(element('w:pStyle', val=f'Heading{level + 1}'))
            lvl.append(element('w:lvlText', val=('•' if kind == 'bullet' else ('%1.' if level == 0 else '%1.%2.'))))
            lvl.append(element('w:lvlJc', val='left'))
            props = element('w:pPr')
            props.append(element('w:ind', left=0 if kind == 'heading' else 310, hanging=0 if kind == 'heading' else 240))
            lvl.append(props)
            abstract.append(lvl)
        root.append(abstract)
        num = element('w:num', numId=num_id)
        num.append(element('w:abstractNumId', val=abstract_id))
        root.append(num)
        return num_id

    def _list_item(self, text: str, num_id: int, level: int = 0):
        """Append a paragraph using the supplied native numbering instance."""
        paragraph = self.doc.add_paragraph()
        props = paragraph._p.get_or_add_pPr()
        num = element('w:numPr')
        num.append(element('w:ilvl', val=level))
        num.append(element('w:numId', val=num_id))
        props.append(num)
        self.inline(paragraph, text)
        return paragraph

    def _field(self, paragraph, instruction: str, cached: str):
        """Append an updateable Word field with a useful cached result."""
        run = paragraph.add_run()
        run._r.append(element('w:fldChar', fldCharType='begin'))
        instr = element('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = ' ' + instruction + ' '
        run._r.append(instr)
        run._r.append(element('w:fldChar', fldCharType='separate'))
        paragraph.add_run(cached)
        paragraph.add_run()._r.append(element('w:fldChar', fldCharType='end'))

    def inline(self, paragraph, text: str):
        """Preserve emphasis and hyperlinks in supported Markdown text."""
        parts = re.split(r'(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\)|`[^`]+`)', text)
        for part in parts:
            if not part:
                continue
            match = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', part)
            if match and match[2].startswith(('https://', 'http://')):
                from docx.opc.constants import RELATIONSHIP_TYPE
                link = OxmlElement('w:hyperlink')
                link.set(qn('r:id'), self.doc.part.relate_to(match[2], RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                run = element('w:r')
                properties = element('w:rPr')
                properties.append(element('w:color', val='00296B'))
                properties.append(element('w:u', val='single'))
                run.append(properties)
                node = element('w:t')
                node.text = match[1]
                run.append(node)
                link.append(run)
                paragraph._p.append(link)
            else:
                run = paragraph.add_run(match[1] if match else part.strip('`') if part.startswith('`') else part[2:-2] if part.startswith('**') else part)
                run.bold = part.startswith('**')

    def cover(self, title: str, subtitle: str):
        """Create a concise identifiable cover without manual blank spacing."""
        p = self.doc.add_paragraph('INSERTEC', 'Subtitle')
        p.paragraph_format.space_before = Pt(75)
        self.doc.add_paragraph(title, 'Title')
        self.doc.add_paragraph(subtitle, 'Subtitle')
        self.doc.add_paragraph(f"Edición documental {self.config['version']} · Entregable {self.revision}")
        months = ('enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio',
                  'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre')
        date = self.publication_date
        self.doc.add_paragraph(f'{date.day} de {months[date.month - 1]} de {date.year} · Español')
        self.callout('DOCUMENTO PARA REVISIÓN', 'Uso interno. Los ejemplos de redacción son ilustrativos. Las acciones disponibles dependen de la compañía, del registro y de los permisos de cada usuario.')
        self.doc.add_paragraph('Responsable editorial: INSERTEC · Documentación CRM')
        self.doc.add_page_break()
        self.doc.add_paragraph('Contenido', 'Title')
        self.toc_container = self.doc.add_paragraph()
        self.doc.add_page_break()

    def heading(self, text: str, level: int):
        """Add a semantic heading and an internal navigation bookmark."""
        p = self.doc.add_heading(text, level)
        if level < 3:
            num = element('w:numPr')
            num.append(element('w:ilvl', val=level - 1))
            num.append(element('w:numId', val=self.heading_num))
            p._p.get_or_add_pPr().append(num)
            if level == 1:
                self.chapter += 1
                self.topic = 0
                if self.chapter > 1:
                    p.paragraph_format.page_break_before = True
            else:
                self.topic += 1
            anchor = f'crm_{len(self.entries) + 1:03d}'
            index = len(self.entries) + 1
            start = element('w:bookmarkStart', id=index, name=anchor)
            end = element('w:bookmarkEnd', id=index)
            p._p.insert(1, start)
            p._p.append(end)
            label = f'{self.chapter}' if level == 1 else f'{self.chapter}.{self.topic}'
            self.entries.append({'label': label, 'text': text, 'level': level, 'anchor': anchor})
        return p

    def callout(self, title: str, text: str):
        """Use paragraph borders for callouts rather than layout tables."""
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.keep_together = True
        props = p._p.get_or_add_pPr()
        border = element('w:pBdr')
        border.append(element('w:left', val='single', sz=18, space=8, color='244C82'))
        props.append(border)
        props.append(element('w:shd', fill='F1F5F9'))
        p.add_run(title + ': ').bold = True
        self.inline(p, text)

    def table(self, rows: list[list[str]]):
        """Create a simple data table with repeated accessible header rows."""
        if len(rows) < 2:
            return
        columns = max(len(row) for row in rows)
        table = self.doc.add_table(rows=0, cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'
        table.autofit = False
        widths = [4.1, 6.3, 6.6] if columns == 3 else [17 / columns] * columns
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        for index, values in enumerate(rows):
            row = table.add_row()
            if index == 0:
                row._tr.get_or_add_trPr().append(element('w:tblHeader'))
            row._tr.get_or_add_trPr().append(element('w:cantSplit'))
            for col, cell in enumerate(row.cells):
                cell.width = Cm(widths[col])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(4)
                self.inline(p, values[col] if col < len(values) else '')
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    if index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor.from_string('FFFFFF')
                if index == 0:
                    cell._tc.get_or_add_tcPr().append(element('w:shd', fill='00296B'))
                elif index % 2 == 0:
                    cell._tc.get_or_add_tcPr().append(element('w:shd', fill='F8FAFC'))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def markdown(self, content: str, topic: bool = False):
        """Render the documented Markdown subset without HTML conversion."""
        content = re.sub(r'<!--.*?-->', '', content, flags=re.S)
        lines = content.splitlines()
        index = 0
        list_id, list_kind = None, None
        while index < len(lines):
            line = lines[index].strip()
            index += 1
            if not line:
                continue
            heading = re.match(r'^(#{1,3})\s+(.+)$', line)
            if heading:
                list_kind = None
                if topic and len(heading[1]) == 1:
                    continue
                self.heading(heading[2], len(heading[1]) + (1 if topic else 0))
                continue
            if line.startswith('|') and index < len(lines) and re.match(r'^\s*\|?\s*:?-{3,}', lines[index]):
                rows = [[v.strip() for v in line.strip('|').split('|')]]
                index += 1
                while index < len(lines) and lines[index].strip().startswith('|'):
                    rows.append([v.strip() for v in lines[index].strip().strip('|').split('|')])
                    index += 1
                self.table(rows)
                list_kind = None
                continue
            numbered = re.match(r'^\d+\.\s+(.*)', line)
            bullet = re.match(r'^[-*•]\s+(.*)', line)
            if numbered or bullet:
                kind = 'ordered' if numbered else 'bullet'
                if kind != list_kind or (numbered and line.startswith('1. ')):
                    list_id = self._number_definition(kind)
                    list_kind = kind
                self._list_item((numbered or bullet)[1], list_id)
                continue
            list_kind = None
            if line.startswith('>'):
                self.callout('Nota', line.lstrip('> '))
                continue
            if line.startswith('!['):
                continue
            self.inline(self.doc.add_paragraph(), line)

    def screenshot(self, data: bytes, descriptor: dict):
        """Insert reviewed original image bytes inline with a caption and alt text."""
        self.figure += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        picture = p.add_run().add_picture(io.BytesIO(data), width=Cm(descriptor.get('widthCm', 7.5)))
        picture._inline.docPr.set('descr', descriptor['alt'])
        picture._inline.docPr.set('title', descriptor['caption'])
        caption = self.doc.add_paragraph(f"Figura {self.figure}. {descriptor['caption']}", 'Caption')
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.media.append({'figure': self.figure, 'sourcePart': descriptor['part'], 'sha256': sha256(data), 'caption': descriptor['caption']})

    def _toc(self):
        """Populate a genuine TOC field with clickable entries and no guessed pages."""
        first = self.toc_container
        first.add_run()._r.append(element('w:fldChar', fldCharType='begin'))
        instr = element('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = ' TOC \\o "1-2" \\h \\z \\u '
        first.add_run()._r.append(instr)
        first.add_run()._r.append(element('w:fldChar', fldCharType='separate'))
        previous = first._p
        for entry in self.entries:
            p = OxmlElement('w:p')
            props = element('w:pPr')
            props.append(element('w:spacing', after=50))
            props.append(element('w:ind', left=200 if entry['level'] == 2 else 0))
            p.append(props)
            link = OxmlElement('w:hyperlink')
            link.set(qn('w:anchor'), entry['anchor'])
            run = element('w:r')
            rpr = element('w:rPr')
            rpr.append(element('w:sz', val=19))
            if entry['level'] == 1:
                rpr.append(element('w:b'))
            run.append(rpr)
            text = element('w:t')
            text.text = entry['label'] + '. ' + entry['text']
            run.append(text)
            link.append(run)
            p.append(link)
            previous.addnext(p)
            previous = p
        end = element('w:p')
        run = element('w:r')
        run.append(element('w:fldChar', fldCharType='end'))
        end.append(run)
        previous.addnext(end)

    def save(self, path: Path, baseline: Path, frozen: list[str]):
        """Save without overwriting and validate authored text, XML and media."""
        if path.exists():
            raise FileExistsError(f'Choose a new version: {path}')
        self._toc()
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            if len(names) != len(set(names)):
                raise ValueError('Duplicate package entries')
            for name in names:
                if name.endswith(('.xml', '.rels')):
                    etree.fromstring(archive.read(name))
            xml = etree.fromstring(archive.read('word/document.xml'))
            paragraphs = [''.join(p.itertext()) for p in xml.findall('.//w:p', self.ns)]
            for text in frozen:
                if text not in paragraphs:
                    raise ValueError('Protected baseline paragraph was not preserved')
            image_hashes = {sha256(archive.read(name)) for name in names if name.startswith('word/media/')}
            if not all(image['sha256'] in image_hashes for image in self.media):
                raise ValueError('An original image was changed')
        manifest = {'document': path.name, 'version': self.config['version'], 'revision': self.revision,
                    'baseline': {'filename': baseline.name, 'sha256': sha256(baseline.read_bytes())},
                    'documentSha256': sha256(path.read_bytes()), 'sourceHashes': self.sources,
                    'headings': self.entries, 'images': self.media,
                    'protectedParagraphs': frozen, 'xmlValid': True,
                    'visualAcceptance': 'pending', 'wordTocPageNumbers': 'refresh in Word; cached navigation has no guessed page numbers'}
        path.with_suffix('.manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def original_currency(baseline: Path) -> list[str]:
    """Read the excluded section verbatim from the user's baseline Word file."""
    with zipfile.ZipFile(baseline) as archive:
        root = etree.fromstring(archive.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    result, active = [], False
    for p in root.findall('.//w:body/w:p', ns):
        text = ''.join(p.xpath('.//w:t/text()', namespaces=ns))
        if text.strip().startswith('8.8.'):
            break
        if active and text.strip():
            result.append(text)
        if text.strip().startswith('8.7.'):
            active = True
    if len(result) != 2:
        raise ValueError('The protected baseline section changed; review its boundaries')
    return result


def main():
    """Resolve canonical inputs and produce the requested cumulative deliverable."""
    parser = argparse.ArgumentParser()
    parser.add_argument('--root', type=Path, default=Path(__file__).resolve().parents[2] / 'docs/crm-help')
    parser.add_argument('--baseline', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--stage', choices=['structure', 'pilots', 'complete'], required=True)
    parser.add_argument('--revision', required=True)
    args = parser.parse_args()
    root = args.root.resolve()
    config = json.loads((root / 'manual/publication.json').read_text(encoding='utf-8-sig'))
    baseline_hash = sha256(args.baseline.read_bytes())
    if baseline_hash != config['baselineSha256']:
        raise ValueError('The baseline does not match the approved July 24 Word file')
    title = {'structure': 'Criterios y estructura del manual CRM', 'pilots': 'Manual CRM · Capítulos piloto', 'complete': 'Manual de uso de la aplicación web CRM'}[args.stage]
    writer = ManualWriter(config, args.revision, title)
    writer.sources['manual/publication.json'] = sha256((root / 'manual/publication.json').read_bytes())
    writer.cover(title, config['subtitles'][args.stage])
    source = root / 'manual' / ('editorial.es-ES.md' if args.stage == 'structure' else 'introduction.es-ES.md')
    writer.sources[source.relative_to(root).as_posix()] = sha256(source.read_bytes())
    writer.markdown(source.read_text(encoding='utf-8-sig'))
    frozen = []
    if args.stage != 'structure':
        with zipfile.ZipFile(args.baseline) as archive:
            for module in config['chapters']:
                if args.stage == 'pilots' and module['id'] not in ('visits', 'expenses'):
                    continue
                writer.heading(module['title'], 1)
                writer.inline(writer.doc.add_paragraph(), module['purpose'])
                for topic_id in module['topics']:
                    if args.stage == 'pilots' and module['id'] == 'expenses' and topic_id not in config['pilotExpenseTopics']:
                        continue
                    folder = root / 'modules' / module['id'] / 'topics' / topic_id
                    metadata_path = folder / 'topic.json'
                    metadata = json.loads(metadata_path.read_text(encoding='utf-8-sig'))
                    writer.sources[metadata_path.relative_to(root).as_posix()] = sha256(metadata_path.read_bytes())
                    path = folder / 'content.es-ES.md'
                    writer.sources[path.relative_to(root).as_posix()] = sha256(path.read_bytes())
                    writer.heading(metadata['title'], 2)
                    if topic_id == 'divisas-e-importes-de-una-linea':
                        frozen = original_currency(args.baseline)
                        for text in frozen:
                            writer.doc.add_paragraph(text)
                    else:
                        writer.markdown(path.read_text(encoding='utf-8-sig'), topic=True)
                    for descriptor in config['screenshots'].get(metadata['id'], []):
                        writer.screenshot(archive.read(descriptor['part']), descriptor)
    writer.save(args.output, args.baseline, frozen)
    if sha256(args.baseline.read_bytes()) != baseline_hash:
        raise ValueError('The baseline changed during generation')
    print(json.dumps({'output': str(args.output), 'headings': len(writer.entries), 'images': len(writer.media), 'protectedParagraphs': len(frozen)}))


if __name__ == '__main__':
    main()
